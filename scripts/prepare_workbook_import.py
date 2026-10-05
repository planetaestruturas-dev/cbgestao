#!/usr/bin/env python3
"""Build a validated CB Gestão import payload from the approved workbooks."""

import json
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


AS_OF = date(2026, 10, 5)
MONTHS = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4,
    "maio": 5, "junho": 6, "julho": 7, "agosto": 8, "setembro": 9,
    "outubro": 10, "novembro": 11, "dezembro": 12,
}


def br_date(value):
    if isinstance(value, datetime):
        value = value.date()
    if not isinstance(value, date):
        raise ValueError(f"Data inválida: {value!r}")
    return value.strftime("%d/%m/%Y")


def parse_money(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").replace("R$", "").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    return float(text)


def normalized(value):
    return "".join(
        char for char in unicodedata.normalize("NFD", str(value or "").lower())
        if unicodedata.category(char) != "Mn"
    ).strip()


def source_rows(ws, start_row):
    for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
        values = list(row)
        if row_number >= start_row and any(value is not None for value in values):
            yield row_number, values


def first_rental_due(start, due_day):
    month = start.month + 1
    year = start.year
    if month == 13:
        month, year = 1, year + 1
    first_complete = date(year, month, min(start.day, 28 if month == 2 else 31))
    last_day = (date(year + (month == 12), month % 12 + 1, 1) - date.resolution).day
    due = date(year, month, min(due_day, last_day))
    if due < first_complete:
        month += 1
        if month == 13:
            month, year = 1, year + 1
        last_day = (date(year + (month == 12), month % 12 + 1, 1) - date.resolution).day
        due = date(year, month, min(due_day, last_day))
    return due


def next_month(value):
    return date(value.year + (value.month == 12), value.month % 12 + 1, 1)


def category_for(description, observation):
    text = normalized(f"{description} {observation}")
    if "reembolso caucao" in text:
        return "Cauções e reembolsos"
    if "tarifa" in text or "anuidade" in text or "estorno" in text or "protesto" in text:
        return "Despesas financeiras › Tarifas e anuidades"
    if "iptu" in text:
        return "Tributos › IPTU"
    if "imposto" in text:
        return "Tributos › Impostos"
    if "reforma" in text or "pintura" in text or "obra" in text or "acal" in text or "leroy" in text:
        return "Obras e reformas › Reforma e material"
    if "detet" in text or "cupim" in text or "fossa" in text or "desentupid" in text:
        return "Serviços e manutenção › Serviços especializados"
    if "analise documentacao" in text:
        return "Serviços e manutenção › Serviços administrativos"
    return "Despesas gerais › Sem classificação"


def main(contracts_file, costs_file, output_file):
    workbook = load_workbook(contracts_file, data_only=True, read_only=True)
    contract_rows = []
    contracts_by_key = {}
    for source_row, values in source_rows(workbook["Contratos"], 4):
        values += [None] * 6
        external_number, unit, tenant, start, end, rent = values[:6]
        if not external_number:
            continue
        start_date = start.date() if isinstance(start, datetime) else start
        end_date = end.date() if isinstance(end, datetime) else end
        external_ref = str(int(external_number)) if isinstance(external_number, float) and external_number.is_integer() else str(external_number)
        # Approved corrections: Livia starts after Manuel, Bianca ends in May,
        # and the prior Kitnet 10 tenancy ends before Shirley begins.
        if unit == "Kitnet 01" and normalized(tenant) == "livia vieira carvalho":
            start_date = date(2026, 6, 20)
        if unit == "Kitnet 06" and "bianca" in normalized(tenant):
            end_date = date(2026, 5, 26)
        if unit == "Kitnet 10" and "maria ludiane" in normalized(tenant):
            end_date = date(2026, 10, 2)
        contract_rows.append({
            "sourceRow": source_row, "externalContractNo": external_ref, "unit": str(unit).strip(),
            "tenant": str(tenant).strip(), "startDate": start_date, "endDate": end_date,
            "rent": parse_money(rent),
        })
        contracts_by_key[(external_ref, normalized(unit), normalized(tenant))] = contract_rows[-1]

    # Payment sheet has a few spelling/combined-name differences.
    def locate_contract(external_ref, unit, tenant):
        candidates = [row for row in contract_rows if row["externalContractNo"] == external_ref and normalized(row["unit"]) == normalized(unit)]
        if external_ref == "47416" and normalized(unit) == "kitnet 01":
            return next(row for row in candidates if "manuel" in normalized(row["tenant"]))
        if external_ref == "5" and normalized(unit) == "kitnet 01":
            return next(row for row in candidates if "livia" in normalized(row["tenant"]))
        if len(candidates) == 1:
            return candidates[0]
        name = normalized(tenant)
        matches = [row for row in candidates if any(part in normalized(row["tenant"]) for part in name.split("/"))]
        if len(matches) == 1:
            return matches[0]
        raise ValueError(f"Contrato não identificado para pagamento {external_ref} / {unit} / {tenant}")

    for index, row in enumerate(contract_rows, 1):
        row["id"] = 1000 + index
        row["dueDay"] = row["startDate"].day
        row["penaltyMultiplier"] = 0
        row["depositAmount"] = 0
        row["depositDue"] = ""
        row["status"] = "Ativo" if row["endDate"] >= AS_OF else "Encerrado"

    payment_rows = []
    for source_row, values in source_rows(workbook["Pagamentos"], 4):
        values += [None] * 11
        external_number, unit, tenant, start, end, rent, paid, due_day, payment_date, status, observation = values[:11]
        if not external_number:
            continue
        external_ref = str(int(external_number)) if isinstance(external_number, float) and external_number.is_integer() else str(external_number)
        payment_date = payment_date.date() if isinstance(payment_date, datetime) else payment_date
        payment_rows.append({
            "sourceRow": source_row, "externalContractNo": external_ref, "unit": str(unit).strip(), "tenant": str(tenant).strip(),
            "rent": parse_money(rent), "paid": parse_money(paid), "dueDay": int(float(due_day)),
            "paymentDate": payment_date, "observation": str(observation or "").strip(),
        })

    charges = []
    expenses = []
    paid_competences = set()
    for payment in payment_rows:
        contract = locate_contract(payment["externalContractNo"], payment["unit"], payment["tenant"])
        # Approved mapping: the 18/06/2026 payment was entered under Livia
        # after the tenancy transfer, but settles Manuel Sena's final rent.
        if payment["sourceRow"] == 12 and payment["externalContractNo"] == "5" and normalized(payment["unit"]) == "kitnet 01":
            contract = next(row for row in contract_rows if row["externalContractNo"] == "47416" and normalized(row["unit"]) == "kitnet 01")
        contract["dueDay"] = payment["dueDay"]
        observation = normalized(payment["observation"])
        if "caucao" in observation:
            contract["depositAmount"] = payment["paid"]
            contract["depositDue"] = br_date(payment["paymentDate"])
            charges.append({
                "id": f"import-deposit-{payment['sourceRow']}", "contractId": contract["id"], "unit": contract["unit"],
                "tenant": contract["tenant"], "due": br_date(payment["paymentDate"]), "competence": "Caução",
                "type": "Caução", "amount": payment["paid"], "status": "Pago", "paidAt": br_date(payment["paymentDate"]),
                "receivedAmount": payment["paid"], "settledAmount": payment["paid"], "depositStatus": "Em posse",
                "receipts": [{"id": f"receipt-{payment['sourceRow']}", "amount": payment["paid"], "receivedOn": br_date(payment["paymentDate"]), "information": payment["observation"], "adjustmentType": "none", "discountAmount": 0, "penaltyInterestAmount": 0, "extraReference": ""}],
            })
            continue
        if "reforma apto 03" in observation:
            expenses.append({
                "id": f"import-expense-livia-{payment['sourceRow']}", "date": br_date(payment["paymentDate"]),
                "dueDate": br_date(payment["paymentDate"]), "paidAt": br_date(payment["paymentDate"]),
                "description": "Reforma apto 03 (pagamento a Livia Vieira Carvalho)", "supplier": "Livia Vieira Carvalho",
                "unit": "Kitnet 03", "maintenanceId": "", "category": "Obras e reformas › Reforma e material",
                "amount": payment["paid"], "status": "Paga",
            })
            continue
        month = next((number for name, number in MONTHS.items() if name in observation), None)
        if month is None:
            raise ValueError(f"Competência de aluguel ausente na linha de pagamento {payment['sourceRow']}: {payment['observation']}")
        # The workbook covers 2026 payment history.
        due_date = date(2026, month, min(payment["dueDay"], (date(2026 + (month == 12), month % 12 + 1, 1) - date.resolution).day))
        competence = due_date.strftime("%Y-%m")
        extra = round(payment["paid"] - contract["rent"], 2)
        charges.append({
            "id": f"import-rent-{payment['sourceRow']}", "contractId": contract["id"], "unit": contract["unit"], "tenant": contract["tenant"],
            "due": br_date(due_date), "competence": competence, "type": "Aluguel", "scheduleSource": "import",
            "amount": contract["rent"], "status": "Pago", "paidAt": br_date(payment["paymentDate"]), "receivedAmount": payment["paid"],
            "settledAmount": contract["rent"], "penaltyInterestAmount": max(0, extra),
            "extraReference": "Multa e juros" if extra > 0 else "", "receiptInformation": payment["observation"],
            "receipts": [{"id": f"receipt-{payment['sourceRow']}", "amount": payment["paid"], "receivedOn": br_date(payment["paymentDate"]), "information": payment["observation"], "adjustmentType": "penaltyInterest" if extra > 0 else "none", "discountAmount": 0, "penaltyInterestAmount": max(0, extra), "extraReference": "Multa e juros" if extra > 0 else ""}],
        })
        paid_competences.add((contract["id"], competence))

    active_contracts = [row for row in contract_rows if row["status"] == "Ativo"]
    for contract in active_contracts:
        first_due = first_rental_due(contract["startDate"], contract["dueDay"])
        cursor = date(max(first_due, AS_OF).year, max(first_due, AS_OF).month, 1)
        while cursor <= contract["endDate"]:
            last_day = (next_month(cursor) - date.resolution).day
            due_date = date(cursor.year, cursor.month, min(contract["dueDay"], last_day))
            competence = due_date.strftime("%Y-%m")
            if due_date >= first_due and (contract["id"], competence) not in paid_competences:
                charges.append({
                    "id": f"contract-{contract['id']}-{competence}", "contractId": contract["id"], "unit": contract["unit"],
                    "tenant": contract["tenant"], "due": br_date(due_date), "competence": competence, "type": "Aluguel",
                    "scheduleSource": "contract", "amount": contract["rent"], "status": "Em aberto", "paidAt": None,
                })
            cursor = next_month(cursor)

    costs_book = load_workbook(costs_file, data_only=True, read_only=True)
    for source_row, values in source_rows(costs_book.active, 3):
        values += [None] * 4
        paid_date, description, amount, observation = values[:4]
        if not isinstance(paid_date, datetime):
            continue
        try:
            parsed_amount = parse_money(amount)
        except (TypeError, ValueError):
            continue
        description = str(description or "").strip()
        observation = str(observation or "").strip()
        expenses.append({
            "id": f"import-expense-cost-{source_row}", "date": br_date(paid_date), "dueDate": br_date(paid_date), "paidAt": br_date(paid_date),
            "description": description + (f" — {observation}" if observation else ""), "supplier": description,
            "unit": "Área comum", "maintenanceId": "", "category": category_for(description, observation), "amount": parsed_amount, "status": "Paga",
        })

    categories = [
        {"id": 1, "name": "Despesas financeiras", "description": "Tarifas, anuidades, estornos e protestos", "parentId": "", "status": "Ativa"},
        {"id": 2, "name": "Tarifas e anuidades", "description": "Custos bancários e de cartão", "parentId": 1, "status": "Ativa"},
        {"id": 3, "name": "Tributos", "description": "Impostos e taxas públicas", "parentId": "", "status": "Ativa"},
        {"id": 4, "name": "IPTU", "description": "Imposto predial e territorial urbano", "parentId": 3, "status": "Ativa"},
        {"id": 5, "name": "Impostos", "description": "Demais impostos", "parentId": 3, "status": "Ativa"},
        {"id": 6, "name": "Obras e reformas", "description": "Obras, melhorias e materiais", "parentId": "", "status": "Ativa"},
        {"id": 7, "name": "Reforma e material", "description": "Reformas, pintura e materiais de obra", "parentId": 6, "status": "Ativa"},
        {"id": 8, "name": "Serviços e manutenção", "description": "Serviços gerais e especializados", "parentId": "", "status": "Ativa"},
        {"id": 9, "name": "Serviços especializados", "description": "Dedetização, fossas e desentupimento", "parentId": 8, "status": "Ativa"},
        {"id": 10, "name": "Serviços administrativos", "description": "Análise documental e serviços administrativos", "parentId": 8, "status": "Ativa"},
        {"id": 11, "name": "Despesas gerais", "description": "Despesas sem classificação específica na origem", "parentId": "", "status": "Ativa"},
        {"id": 12, "name": "Sem classificação", "description": "Origem sem observação de categoria", "parentId": 11, "status": "Ativa"},
        {"id": 13, "name": "Cauções e reembolsos", "description": "Devoluções e ajustes de caução", "parentId": "", "status": "Ativa"},
    ]
    supplier_names = sorted({expense["supplier"] for expense in expenses if expense["supplier"]}, key=normalized)
    suppliers = [{"id": index + 1, "name": name, "document": "", "phone": "", "category": "Importado da planilha", "status": "Ativo"} for index, name in enumerate(supplier_names)]
    tenant_names = []
    for contract in contract_rows:
        if contract["tenant"] not in tenant_names:
            tenant_names.append(contract["tenant"])
    tenants = [{"id": index + 1, "name": name, "cpf": "", "phone": "", "email": "", "status": "Ativo", "observations": "Importado da planilha de contratos.", "attachments": []} for index, name in enumerate(tenant_names)]
    units = []
    for unit_number in range(1, 13):
        unit_name = f"Kitnet {unit_number:02d}"
        current = next(contract for contract in active_contracts if contract["unit"] == unit_name)
        units.append({"id": unit_number, "name": unit_name, "status": "Ocupada", "tenant": current["tenant"], "rent": current["rent"], "description": "", "cagece": "", "enel": "", "iptu": "", "attachments": []})
    contracts = [{
        "id": row["id"], "externalContractNo": row["externalContractNo"], "unit": row["unit"], "tenant": row["tenant"],
        "start": br_date(row["startDate"]), "end": br_date(row["endDate"]), "dueDay": row["dueDay"], "penaltyMultiplier": row["penaltyMultiplier"],
        "rent": row["rent"], "depositAmount": row["depositAmount"], "depositDue": row["depositDue"], "status": row["status"], "attachmentName": "", "inspectionAttachments": [],
    } for row in contract_rows]
    payload = {
        "cb-gestao:unidades": units, "cb-gestao:cobrancas": charges, "cb-gestao:despesas": expenses,
        "cb-gestao:extrato": [], "cb-gestao:contratos": contracts, "cb-gestao:manutencoes": [], "cb-gestao:distratos": [],
        "cb-gestao:inquilinos": tenants,
        "cb-gestao:contas-bancarias": [
            {"id": 1, "bank": "Banco Inter", "account": "Histórico importado", "type": "Conta corrente", "status": "Ativa"},
            {"id": 2, "bank": "Banco Bradesco", "account": "Histórico importado", "type": "Conta corrente", "status": "Ativa"},
        ],
        "cb-gestao:fornecedores": suppliers, "cb-gestao:categorias-despesa": categories, "cb-gestao:pontos-restauracao": [],
    }
    Path(output_file).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"contracts": len(contracts), "tenants": len(tenants), "charges": len(charges), "paidCharges": sum(1 for item in charges if item["status"] == "Pago"), "openCharges": sum(1 for item in charges if item["status"] == "Em aberto"), "expenses": len(expenses), "suppliers": len(suppliers), "expenseTotal": round(sum(item["amount"] for item in expenses), 2)}, ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("Uso: prepare_workbook_import.py contratos.xlsx custos.xlsx saida.json")
    main(*sys.argv[1:])
